"""OEM codes: shown on every part, filled in by staff, sent in by buyers and mechanics."""

from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Product, ReferenceNumber
from app.oem import split_numbers, valid_oem
from app.price_list_import import import_price_list, supplier_by_name
from tests.conftest import admin_login
from tests.test_price_lists import RECEIVED, client, engine, workbook  # noqa: F401

SESSION = "oem-test-session"


def _missing(client) -> dict[str, dict]:  # noqa: F811
    return {row["name"]: row for row in client.get("/api/admin/oem/missing").json()["items"]}


def _relay_id(client) -> str:  # noqa: F811
    return next(row["id"] for name, row in _missing(client).items() if "relesi" in name)


def _oem(client, product_id: str) -> list[str]:  # noqa: F811
    detail = client.get(f"/api/v1/products/{product_id}").json()
    return [ref["number"] for ref in detail["reference_numbers"] if ref["type"] == "OE"]


def test_codes_in_a_cell_are_split_and_checked() -> None:
    assert split_numbers("96267421/22/23") == ["96267421", "96267422", "96267423"]
    assert split_numbers("58101-1RA00, 96 273 708") == ["58101-1RA00", "96273708"]
    assert [valid_oem(code) for code in ("96273708", "LR001234", "0644", "KOLODKA")] == [
        True, True, False, False,
    ]


def test_staff_fill_in_missing_oem_codes(client) -> None:  # noqa: F811
    admin_login(client)
    summary = client.get("/api/admin/oem/summary").json()
    assert summary["total"] - summary["with_oem"] == 2
    missing = _missing(client)
    assert len(missing) == 2
    relay = next(row for name, row in missing.items() if "relesi" in name)
    assert relay["cars"] == ["Daewoo Nexia"]

    rejected = client.post(f"/api/admin/oem/parts/{relay['id']}", json={"numbers": ["0644"]})
    assert rejected.status_code == 422
    saved = client.post(f"/api/admin/oem/parts/{relay['id']}", json={"numbers": ["96 252 427"]})
    assert saved.status_code == 200, saved.text
    assert saved.json()["added"] == ["96252427"]
    assert _oem(client, relay["id"]) == ["96252427"]
    found = client.get("/api/v1/search", params={"q": "96252427"}).json()["items"]
    assert [item["id"] for item in found] == [relay["id"]]
    assert client.get("/api/admin/oem/summary").json()["with_oem"] == summary["with_oem"] + 1

    # A typo is removed again; codes from a seller's list are not removed here.
    removed = client.delete(f"/api/admin/oem/parts/{relay['id']}/96252427")
    assert removed.status_code == 200 and removed.json()["oem"] == []
    shock = client.get("/api/v1/search", params={"q": "96394572"}).json()["items"][0]
    assert client.delete(f"/api/admin/oem/parts/{shock['id']}/96394572").status_code == 404


def test_buyers_send_codes_that_staff_check_first(client) -> None:  # noqa: F811
    admin_login(client)
    relay = _relay_id(client)
    url = f"/api/v1/products/{relay}/oem-suggestions"
    assert client.post(url, json={"number": "kolodka", "session_id": SESSION}).status_code == 422
    for _ in range(2):
        assert client.post(url, json={"number": "96252427", "session_id": SESSION}).status_code == 201
    # Not shown until staff accept it.
    assert _oem(client, relay) == []
    suggestions = client.get("/api/admin/oem/suggestions").json()
    assert [(row["product_id"], row["number"]) for row in suggestions] == [(relay, "96252427")]
    # Staff see it next to the part in the list of parts to fill in.
    row = next(row for name, row in _missing(client).items() if "relesi" in name)
    assert [item["number"] for item in row["suggestions"]] == ["96252427"]

    accepted = client.patch(
        f"/api/admin/oem/suggestions/{suggestions[0]['id']}", json={"status": "ACCEPTED"}
    )
    assert accepted.status_code == 200
    assert _oem(client, relay) == ["96252427"]
    assert client.get("/api/admin/oem/suggestions").json() == []


def test_excel_download_and_upload(client) -> None:  # noqa: F811
    admin_login(client)
    download = client.get("/api/admin/oem/missing.xlsx")
    assert download.status_code == 200
    book = load_workbook(BytesIO(download.content))
    sheet = book["Parts without OEM"]
    assert sheet["A1"].value == "Part ID" and sheet["J1"].value == "OEM code"
    rows = {sheet.cell(row, 2).value: row for row in range(2, sheet.max_row + 1)}
    relay_row = next(row for name, row in rows.items() if "relesi" in name)
    other_row = next(row for name, row in rows.items() if "relesi" not in name)
    sheet.cell(relay_row, 10, "96252427, 96252428")
    sheet.cell(other_row, 10, "STD")
    sheet.append(["P-999999", None, None, None, None, None, None, None, None, "96000000"])
    filled = BytesIO()
    book.save(filled)

    result = client.post(
        "/api/admin/oem/upload",
        files={"file": ("filled.xlsx", filled.getvalue(), "application/octet-stream")},
    ).json()
    relay = sheet.cell(relay_row, 1).value
    assert result["codes_added"] == 2 and result["parts_updated"] == 1
    assert result["unknown_parts"] == ["P-999999"]
    assert result["invalid"] == [f"{sheet.cell(other_row, 1).value}: STD"]
    assert _oem(client, relay) == ["96252427", "96252428"]

    not_excel = client.post(
        "/api/admin/oem/upload", files={"file": ("x.xlsx", b"not a workbook", "text/plain")}
    )
    assert not_excel.status_code == 422


def test_a_sellers_oem_column_is_saved_as_oem(engine) -> None:  # noqa: F811
    content = workbook([
        ["Наименование", "OEM", "Модел"],
        ["РЕЛЕ БЕНЗОНАСОСА", "96252427", "NEXIA"],
        ["FARA OLDI CHAP", "LR001234", "NEXIA"],
    ])
    with Session(engine) as session:
        supplier = supplier_by_name(session, "OEM DO'KON")
        import_price_list(session, supplier, content, "oem.xlsx", RECEIVED)
        session.commit()
        numbers = {
            (name, number, kind)
            for name, number, kind in session.execute(
                select(Product.name, ReferenceNumber.number, ReferenceNumber.type)
                .join(ReferenceNumber, ReferenceNumber.product_id == Product.id)
            )
        }
    # "LR001234" does not look like a GM number, but the seller said it is the OEM code.
    assert {(number, kind) for _, number, kind in numbers} == {
        ("96252427", "OE"), ("LR001234", "OE"),
    }
