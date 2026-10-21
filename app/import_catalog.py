from __future__ import annotations

import argparse
from collections import defaultdict
from collections.abc import Iterable
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from app.catalog import refresh_products
from app.database import SessionLocal
from app.models import (
    Brand,
    CatalogFieldDoc,
    CatalogListValue,
    CatalogReviewItem,
    CatalogSetting,
    Category,
    CategoryAttribute,
    Fitment,
    FitReport,
    KitComponent,
    OfferImportReview,
    PilotFeedback,
    Product,
    ProductAttribute,
    ProductAudit,
    ProductMedia,
    RawSupplierProduct,
    ReferenceNumber,
    ReservationRequest,
    SearchEvent,
    SearchEventAction,
    SearchSynonym,
    Supersession,
    Supplier,
    SupplierImport,
    SupplierOffer,
    SupplierProductMapping,
    Vehicle,
)
from app.normalization import normalize_code, normalize_text
from app.part_types import ensure_standard_part_types

AUDIT_SEVERITY = {"CONFLICT": 3, "REVIEW": 2, "PROBABLE": 1}


def _text(value: Any) -> str | None:
    if value is None:
        return None
    result = str(value).strip()
    return result or None


def _integer(value: Any) -> int | None:
    if value in (None, ""):
        return None
    return int(value)


def _decimal(value: Any) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError(f"Invalid decimal value: {value!r}") from exc


def _rows(worksheet: Any) -> Iterable[dict[str, Any]]:
    iterator = worksheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(iterator)]
    for values in iterator:
        row = {header: value for header, value in zip(headers, values) if header}
        if any(value not in (None, "") for value in row.values()):
            yield row


def _v2_rows(worksheet: Any) -> Iterable[dict[str, Any]]:
    iterator = worksheet.iter_rows(values_only=True)
    headers = [
        str(value or "").replace(" *", "").strip().casefold().replace(" ", "_")
        for value in next(iterator)
    ]
    for values in iterator:
        row = {header: value for header, value in zip(headers, values) if header}
        if any(value not in (None, "") for value in row.values()):
            yield row


def _clear_catalog(session: Session) -> None:
    for model in (
        PilotFeedback,
        FitReport,
        ReservationRequest,
        SearchEventAction,
        SearchEvent,
        OfferImportReview,
        SupplierImport,
        SupplierProductMapping,
        RawSupplierProduct,
        CatalogReviewItem,
        ProductAudit,
        ProductMedia,
        KitComponent,
        Supersession,
        ProductAttribute,
        SupplierOffer,
        Fitment,
        ReferenceNumber,
        SearchSynonym,
        Vehicle,
        Product,
        CategoryAttribute,
        Category,
        Brand,
        Supplier,
        CatalogSetting,
        CatalogListValue,
        CatalogFieldDoc,
    ):
        session.execute(delete(model))


def _flag(value: Any, default: bool = False) -> bool:
    text = _text(value)
    if text is None:
        return default
    return text.upper() in {"Y", "YES", "TRUE", "1", "HA"}


def _whole(value: Any) -> int | None:
    """Integer cells may arrive as int, float, or text ("1996", "4.0")."""
    if value in (None, ""):
        return None
    try:
        return int(Decimal(str(value).strip()))
    except (InvalidOperation, ValueError):
        return None


def _number(value: Any) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value).strip().replace(",", "."))
    except InvalidOperation:
        return None


def _moment(value: Any) -> datetime | None:
    return value if isinstance(value, datetime) else None


def _optional_sheet(workbook: Any, name: str) -> list[dict[str, Any]]:
    return list(_v2_rows(workbook[name])) if name in workbook.sheetnames else []


def _list_values(workbook: Any) -> list[CatalogListValue]:
    if "Lists" not in workbook.sheetnames:
        return []
    iterator = workbook["Lists"].iter_rows(values_only=True)
    names = [_text(value) for value in next(iterator)]
    columns: dict[str, list[str]] = defaultdict(list)
    for values in iterator:
        for name, value in zip(names, values):
            text = _text(value)
            if name and text and text not in columns[name]:
                columns[name].append(text)
    return [
        CatalogListValue(list_name=name, value=value, sort_order=order)
        for name, values in columns.items()
        for order, value in enumerate(values)
    ]


def _settings(workbook: Any) -> list[CatalogSetting]:
    """The Settings sheet holds key/value rows plus a currency -> rate_to_uzs table."""
    if "Settings" not in workbook.sheetnames:
        return []
    settings: dict[str, CatalogSetting] = {}
    currency_table = False
    for key, value, note, *_ in workbook["Settings"].iter_rows(min_row=2, values_only=True):
        key_text = _text(key)
        if key_text is None:
            currency_table = False
            continue
        if key_text == "currency" and _text(value) == "rate_to_uzs":
            currency_table = True
            continue
        name = f"rate_to_uzs.{key_text}" if currency_table else key_text
        stored = value.isoformat() if isinstance(value, datetime) else _text(value)
        settings[name] = CatalogSetting(key=name, value=stored, note=_text(note))
    return list(settings.values())


def _field_docs(workbook: Any) -> list[CatalogFieldDoc]:
    docs: dict[tuple[str, str], CatalogFieldDoc] = {}
    for row in _optional_sheet(workbook, "Data Dictionary"):
        sheet, column = _text(row.get("sheet")), _text(row.get("column"))
        if sheet and column:
            docs[(sheet, column)] = CatalogFieldDoc(
                sheet=sheet,
                column=column,
                required=_flag(row.get("required")),
                kind=_text(row.get("kind")),
                allowed_values=_text(row.get("allowed_values_/_link")),
                description=_text(row.get("description")),
            )
    return list(docs.values())


def _import_v2_workbook(session: Session, workbook: Any, workbook_path: Path) -> dict[str, int]:
    required = {
        "Brands",
        "Products",
        "References",
        "Vehicles",
        "Fitments",
        "Suppliers",
        "Supplier Offers",
        "Raw Mapping",
        "Cross Supplier Audit",
    }
    missing = required.difference(workbook.sheetnames)
    if missing:
        raise ValueError(f"V2 workbook is missing sheets: {', '.join(sorted(missing))}")

    # Reference data: dropdown lists, help text, and settings drive the catalog editor.
    list_values = _list_values(workbook)
    field_docs = _field_docs(workbook)
    settings = _settings(workbook)
    session.add_all([*list_values, *field_docs, *settings])

    brands_by_source_id: dict[str, Brand] = {}
    for row in _v2_rows(workbook["Brands"]):
        source_id = _text(row.get("brand_id"))
        name = _text(row.get("brand_name"))
        if not source_id or not name:
            continue
        brands_by_source_id[source_id] = Brand(
            name=name,
            normalized_name=normalize_text(name),
            source_id=source_id,
            code=_text(row.get("brand_code")),
            aliases=_text(row.get("aliases")),
            brand_type=_text(row.get("brand_type")),
            quality_tier=_text(row.get("quality_tier")),
            country=_text(row.get("country")),
            tecdoc_brand_id=_text(row.get("tecdoc_brand_id")),
            website=_text(row.get("website")),
            notes=_text(row.get("notes")),
        )
    session.add_all(brands_by_source_id.values())

    categories: dict[str, Category] = {}
    category_rows = _optional_sheet(workbook, "Categories")
    for row in category_rows:
        category_id = _text(row.get("category_id"))
        name = _text(row.get("name_uz")) or _text(row.get("name_ru"))
        if not category_id or not name:
            continue
        categories[category_id] = Category(
            id=category_id,
            name_uz=name,
            name_ru=_text(row.get("name_ru")),
            name_en=_text(row.get("name_en")),
            is_leaf=_flag(row.get("is_leaf"), default=True),
            tecdoc_generic_article_id=_text(row.get("tecdoc_generic_article_id")),
            pcdb_part_type_id=_text(row.get("pcdb_part_type_id")),
        )
    session.add_all(categories.values())
    session.flush()
    for row in category_rows:
        category = categories.get(_text(row.get("category_id")) or "")
        parent_id = _text(row.get("parent_id"))
        if category and parent_id in categories:
            category.parent_id = parent_id

    category_attributes: list[CategoryAttribute] = []
    category_attribute_seen: set[tuple[str, str]] = set()
    for row in _optional_sheet(workbook, "Category Attributes"):
        category_id = _text(row.get("category_id"))
        code = _text(row.get("attribute_code"))
        if category_id not in categories or not code:
            continue
        if (category_id, code) in category_attribute_seen:
            continue
        category_attribute_seen.add((category_id, code))
        category_attributes.append(
            CategoryAttribute(
                category_id=category_id,
                attribute_code=code,
                name_uz=_text(row.get("name_uz")) or code,
                name_ru=_text(row.get("name_ru")),
                unit=_text(row.get("unit")),
                data_type=(_text(row.get("data_type")) or "TEXT").upper(),
                required=_flag(row.get("required")),
                allowed_values=_text(row.get("allowed_values")),
                is_filter=_flag(row.get("is_filter")),
                sort_order=_whole(row.get("sort_order")),
            )
        )
    session.add_all(category_attributes)

    suppliers_by_source_id: dict[str, Supplier] = {}
    for row in _v2_rows(workbook["Suppliers"]):
        source_id = _text(row.get("supplier_id"))
        name = _text(row.get("supplier_name"))
        if not source_id or not name:
            continue
        suppliers_by_source_id[source_id] = Supplier(
            name=name,
            normalized_name=normalize_text(name),
            source_id=source_id,
            supplier_type=_text(row.get("supplier_type")),
            city=_text(row.get("city")),
            address=_text(row.get("city")),
            contact=_text(row.get("contact")),
            default_currency=(_text(row.get("default_currency")) or "UZS")[:3].upper(),
            integration=_text(row.get("integration")),
            update_frequency_days=_whole(row.get("update_frequency_days")),
            default_lead_time_days=_whole(row.get("default_lead_time_days")),
            reliability_score=_number(row.get("reliability_score")),
            is_active=_flag(row.get("active"), default=True),
            created_at=datetime.now(UTC),
        )
    session.add_all(suppliers_by_source_id.values())
    session.flush()

    audit_rows = [
        row for row in _v2_rows(workbook["Cross Supplier Audit"]) if _text(row.get("product_id"))
    ]
    audits_by_product: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in audit_rows:
        audits_by_product[_text(row.get("product_id"))].append(row)

    products: dict[str, Product] = {}
    for row in _v2_rows(workbook["Products"]):
        product_id = _text(row.get("product_id"))
        name = _text(row.get("product_name_uz")) or _text(row.get("product_name_ru"))
        if not product_id or not name:
            continue
        findings = audits_by_product.get(product_id, [])
        # The most serious finding decides the product's status (CONFLICT hides it).
        worst = max(
            findings,
            key=lambda finding: AUDIT_SEVERITY.get(
                (_text(finding.get("merge_status")) or "").upper(), 0
            ),
            default={},
        )
        status = (_text(row.get("status")) or "ACTIVE").upper()
        products[product_id] = Product(
            id=product_id,
            name=name,
            normalized_name=normalize_text(name),
            name_ru=_text(row.get("product_name_ru")),
            name_en=_text(row.get("product_name_en")),
            brand=brands_by_source_id.get(_text(row.get("brand_id"))),
            category_id=_text(row.get("category_id"))
            if _text(row.get("category_id")) in categories
            else None,
            article_number=_text(row.get("article_number")),
            origin=_text(row.get("origin")),
            status=status,
            unit=_text(row.get("unit")),
            pack_qty=_whole(row.get("pack_qty")),
            weight_kg=_number(row.get("weight_kg")),
            length_mm=_number(row.get("length_mm")),
            width_mm=_number(row.get("width_mm")),
            height_mm=_number(row.get("height_mm")),
            ean=_text(row.get("ean")),
            country_of_origin=_text(row.get("country_of_origin")),
            description=_text(row.get("description")),
            universal=_flag(row.get("universal")),
            source_rows=_text(row.get("source_rows")),
            merge_basis=_text(row.get("merge_basis")),
            audit_status=(_text(worst.get("merge_status")) or status).upper(),
            audit_reason=" || ".join(
                filter(None, (_text(finding.get("audit_reason")) for finding in findings))
            )
            or None,
        )
    session.add_all(products.values())
    # Audits, review items, kits, and media reference products by column only, so the
    # unit of work cannot order them after products on its own.
    session.flush()

    attributes: list[ProductAttribute] = []
    attribute_seen: set[tuple[str, str]] = set()
    for row in _optional_sheet(workbook, "Attributes"):
        product_id = _text(row.get("product_id"))
        code = _text(row.get("attribute_code"))
        value = _text(row.get("value"))
        if product_id not in products or not code or value is None:
            continue
        if (product_id, code) in attribute_seen:
            continue
        attribute_seen.add((product_id, code))
        attributes.append(
            ProductAttribute(
                product_id=product_id,
                attribute_code=code,
                value=value,
                unit=_text(row.get("unit")),
                source=_text(row.get("source")),
            )
        )
    session.add_all(attributes)

    audits = [
        ProductAudit(
            product_id=product_id,
            brand=_text(row.get("brand")),
            reference=_text(row.get("reference")),
            master_name=_text(row.get("master_name")),
            supplier_count=_whole(row.get("supplier_count")),
            suppliers=_text(row.get("suppliers")),
            cross_name_similarity=_text(row.get("cross_name_similarity")),
            internal_collision=_text(row.get("internal_collision")),
            positions=_text(row.get("positions")),
            sides=_text(row.get("sides")),
            merge_status=(_text(row.get("merge_status")) or "UNAUDITED").upper(),
            audit_reason=_text(row.get("audit_reason")),
            source_rows=_text(row.get("source_rows")),
            attr_conflict=_text(row.get("attr_conflict")),
        )
        for row in audit_rows
        if (product_id := _text(row.get("product_id"))) in products
    ]
    session.add_all(audits)

    synonym_rows = _optional_sheet(workbook, "Search Synonyms")
    synonyms_by_vehicle: dict[str, list[str]] = defaultdict(list)
    synonyms: list[SearchSynonym] = []
    for row in synonym_rows:
        synonym_id = _text(row.get("synonym_id"))
        target_type = (_text(row.get("target_type")) or "").upper()
        target_id = _text(row.get("target_id"))
        term = _text(row.get("term"))
        if not synonym_id or not target_type or not target_id or not term:
            continue
        synonyms.append(
            SearchSynonym(
                id=synonym_id,
                target_type=target_type,
                target_id=target_id,
                term=term,
                lang=_text(row.get("lang")),
                synonym_type=_text(row.get("synonym_type")),
            )
        )
        if target_type == "VEHICLE":
            synonyms_by_vehicle[target_id].append(term)
    session.add_all(synonyms)

    vehicles: dict[str, Vehicle] = {}
    for row in _v2_rows(workbook["Vehicles"]):
        vehicle_id = _text(row.get("vehicle_id"))
        model = _text(row.get("model"))
        if not vehicle_id or not model:
            continue
        make = _text(row.get("make"))
        generation = _text(row.get("generation"))
        modification = _text(row.get("modification"))
        aliases = _text(row.get("aliases"))
        searchable = [
            make,
            model,
            generation,
            modification,
            aliases,
            _text(row.get("display_name")),
            *synonyms_by_vehicle[vehicle_id],
        ]
        vehicles[vehicle_id] = Vehicle(
            id=vehicle_id,
            make=make,
            model=model,
            generation=" · ".join(filter(None, [generation, modification])) or None,
            normalized_text=normalize_text(" ".join(filter(None, searchable)))[:600],
            generation_code=generation,
            modification=modification,
            year_from=_whole(row.get("year_from")),
            year_to=_whole(row.get("year_to")),
            body_type=_text(row.get("body_type")),
            engine_code=_text(row.get("engine_code")),
            displacement_cc=_whole(row.get("displacement_cc")),
            power_kw=_whole(row.get("power_kw")),
            fuel_type=_text(row.get("fuel_type")),
            transmission=_text(row.get("transmission")),
            drive_type=_text(row.get("drive_type")),
            market=_text(row.get("market")),
            vin_wmi=_text(row.get("vin_wmi")),
            vin_vds_pattern=_text(row.get("vin_vds_pattern")),
            tecdoc_ktype=_text(row.get("tecdoc_ktype")),
            vehicle_source=_text(row.get("vehicle_source")),
            origin_type=_text(row.get("origin_type")),
            plant=_text(row.get("plant")),
            aliases=aliases,
            parc_priority=_text(row.get("parc_priority")),
            data_confidence=_text(row.get("data_confidence")),
            notes=_text(row.get("notes")),
        )
    session.add_all(vehicles.values())
    session.flush()

    references: list[ReferenceNumber] = []
    reference_seen: set[tuple[str, str, str]] = set()
    skipped_supplier_codes = 0
    for row in _v2_rows(workbook["References"]):
        product_id = _text(row.get("product_id"))
        number = _text(row.get("number_raw"))
        reference_type = (_text(row.get("ref_type")) or "UNKNOWN").upper()
        if not product_id or product_id not in products or not number:
            continue
        if reference_type == "SUPPLIER_INTERNAL_CODE":
            skipped_supplier_codes += 1
            continue
        normalized_number = normalize_code(number)
        key = (product_id, normalized_number, reference_type)
        if not normalized_number or key in reference_seen:
            continue
        reference_seen.add(key)
        references.append(
            ReferenceNumber(
                product_id=product_id,
                number=number,
                normalized_number=normalized_number,
                type=reference_type,
                source_supplier=suppliers_by_source_id.get(_text(row.get("source_supplier_id"))),
                source_id=_text(row.get("ref_id")),
                number_brand=brands_by_source_id.get(_text(row.get("number_brand_id"))),
                cross_source=_text(row.get("cross_source")),
                confidence=_text(row.get("confidence")),
                verified=_flag(row.get("verified")),
                verified_by=_text(row.get("verified_by")),
                verified_at=_moment(row.get("verified_at")),
                notes=_text(row.get("notes")),
            )
        )
    session.add_all(references)

    fitments: list[Fitment] = []
    fitment_seen: set[tuple[str, str, str | None]] = set()
    for row in _v2_rows(workbook["Fitments"]):
        product_id = _text(row.get("product_id"))
        vehicle_id = _text(row.get("vehicle_id"))
        source_supplier_id = _text(row.get("source_supplier_id"))
        key = (product_id or "", vehicle_id or "", source_supplier_id)
        if (
            not product_id
            or product_id not in products
            or not vehicle_id
            or vehicle_id not in vehicles
            or key in fitment_seen
        ):
            continue
        fitment_seen.add(key)
        fitments.append(
            Fitment(
                product_id=product_id,
                vehicle_id=vehicle_id,
                source_vehicle_text=_text(row.get("source_vehicle_text")),
                source_supplier=suppliers_by_source_id.get(source_supplier_id),
                source_id=_text(row.get("fitment_id")),
                position=_text(row.get("position")),
                side=_text(row.get("side")),
                vertical=_text(row.get("vertical")),
                qty_per_vehicle=_whole(row.get("qty_per_vehicle")),
                year_from=_whole(row.get("year_from")),
                year_to=_whole(row.get("year_to")),
                vin_from=_text(row.get("vin_from")),
                vin_to=_text(row.get("vin_to")),
                condition_note=_text(row.get("condition_note")),
                fitment_source=_text(row.get("fitment_source")),
                confidence=_text(row.get("confidence")),
                verified=_flag(row.get("verified")),
            )
        )
    session.add_all(fitments)

    offers: list[SupplierOffer] = []
    active_offer_keys: set[tuple[str, str, str]] = set()
    for row in _v2_rows(workbook["Supplier Offers"]):
        product_id = _text(row.get("product_id"))
        source_supplier_id = _text(row.get("supplier_id"))
        supplier = suppliers_by_source_id.get(source_supplier_id)
        source_offer_id = _text(row.get("offer_id"))
        if not product_id or product_id not in products or not supplier or not source_offer_id:
            continue
        sku = _text(row.get("supplier_sku_or_ref"))
        normalized_sku = normalize_code(sku) or None
        offer_key = (source_supplier_id or "", product_id, normalized_sku or source_offer_id)
        is_active = offer_key not in active_offer_keys and supplier.is_active
        active_offer_keys.add(offer_key)
        offers.append(
            SupplierOffer(
                source_offer_id=source_offer_id,
                supplier=supplier,
                product_id=product_id,
                supplier_sku=sku,
                normalized_supplier_sku=normalized_sku,
                price=_decimal(row.get("price_uzs") or row.get("price")),
                currency=(_text(row.get("currency")) or "UZS")[:3].upper(),
                stock_qty=_integer(row.get("stock_qty")),
                warehouse=_text(row.get("warehouse")),
                updated_at=_moment(row.get("updated_at")),
                source_type="EXCEL",
                source_reference=f"{workbook_path.name}:{source_offer_id}",
                is_active=is_active,
                data_status="DEMO",
                price_type=_text(row.get("price_type")),
                lead_time_days=_whole(row.get("lead_time_days")),
                min_order_qty=_whole(row.get("min_order_qty")),
                condition=_text(row.get("condition")),
                source_data_status=_text(row.get("data_status")),
            )
        )
    session.add_all(offers)

    supersessions = [
        Supersession(
            old_number_raw=_text(row.get("old_number_raw")),
            old_brand_id=getattr(
                brands_by_source_id.get(_text(row.get("old_brand_id")) or ""), "id", None
            ),
            old_product_id=_text(row.get("old_product_id"))
            if _text(row.get("old_product_id")) in products
            else None,
            new_product_id=_text(row.get("new_product_id")),
            interchangeability=_text(row.get("interchangeability")),
            effective_date=_moment(row.get("effective_date")),
            source=_text(row.get("source")),
            notes=_text(row.get("notes")),
        )
        for row in _optional_sheet(workbook, "Supersessions")
        if _text(row.get("old_number_raw")) and _text(row.get("new_product_id")) in products
    ]
    kits = [
        KitComponent(
            kit_product_id=_text(row.get("kit_product_id")),
            component_product_id=_text(row.get("component_product_id")),
            qty=_whole(row.get("qty")) or 1,
            notes=_text(row.get("notes")),
        )
        for row in _optional_sheet(workbook, "Kits")
        if _text(row.get("kit_product_id")) in products
        and _text(row.get("component_product_id")) in products
    ]
    media = [
        ProductMedia(
            id=_text(row.get("media_id")),
            product_id=_text(row.get("product_id")),
            media_type=_text(row.get("media_type")) or "IMAGE",
            url_or_path=_text(row.get("url_or_path")),
            is_primary=_flag(row.get("is_primary")),
            source=_text(row.get("source")),
            license_note=_text(row.get("license_note")),
        )
        for row in _optional_sheet(workbook, "Media")
        if _text(row.get("media_id"))
        and _text(row.get("product_id")) in products
        and _text(row.get("url_or_path"))
    ]
    session.add_all([*supersessions, *kits, *media])

    review_rows = _optional_sheet(workbook, "Review Queue")
    review_items = [
        CatalogReviewItem(
            raw_id=_text(row.get("raw_id")),
            supplier=_text(row.get("supplier")),
            product_id=_text(row.get("product_id"))
            if _text(row.get("product_id")) in products
            else None,
            raw_name=_text(row.get("raw_name")),
            parsed_product_name=_text(row.get("parsed_product_name")),
            brand=_text(row.get("brand")),
            reference=_text(row.get("reference")),
            vehicle_raw=_text(row.get("vehicle_raw")),
            review_reason=_text(row.get("review_reason")) or "",
            status=(_text(row.get("status")) or "TO REVIEW").upper(),
            reviewer=_text(row.get("reviewer")),
            reviewed_at=_moment(row.get("reviewed_at")),
            resolution_note=_text(row.get("resolution_note")),
        )
        for row in review_rows
        if _text(row.get("raw_id"))
    ]
    session.add_all(review_items)

    review_by_raw = {item.raw_id: item for item in review_items}
    raw_products: list[RawSupplierProduct] = []
    for row in _v2_rows(workbook["Raw Mapping"]):
        raw_id = _text(row.get("raw_id"))
        product_id = _text(row.get("product_id"))
        supplier = suppliers_by_source_id.get(_text(row.get("supplier_id")))
        raw_name = _text(row.get("raw_name"))
        if not raw_id or not supplier or not raw_name:
            continue
        review = review_by_raw.get(raw_id)
        audit_status = (
            "CONFLICT"
            if product_id in products and products[product_id].audit_status == "CONFLICT"
            else (
                (review.status if review else None) or _text(row.get("match_status")) or "UNAUDITED"
            )
        )
        raw_products.append(
            RawSupplierProduct(
                source_raw_id=raw_id,
                supplier=supplier,
                source_row=_integer(row.get("source_row")),
                product_id=product_id if product_id in products else None,
                raw_name=raw_name,
                parsed_product_name=_text(row.get("parsed_product_name")),
                raw_brand=_text(row.get("brand")),
                supplier_code=_text(row.get("reference")),
                supplier_code_type=_text(row.get("reference_type")),
                origin=_text(row.get("origin")),
                unit=_text(row.get("unit")),
                vehicle_raw=_text(row.get("vehicle_raw")),
                position=_text(row.get("position")),
                side=_text(row.get("side")),
                engine_hint=_text(row.get("engine_hint")),
                audit_status=audit_status,
                audit_reason=review.review_reason if review else None,
                source_payload={
                    key: value.isoformat() if isinstance(value, datetime) else value
                    for key, value in row.items()
                    if value is not None
                },
            )
        )
    session.add_all(raw_products)

    session.flush()
    standard_part_types = ensure_standard_part_types(session)
    session.expire_all()
    refresh_products(session)
    session.commit()
    return {
        "products": len(products),
        "quarantined_conflicts": sum(1 for audit in audits if audit.merge_status == "CONFLICT"),
        "brands": len(brands_by_source_id),
        "categories": len(categories),
        "standard_part_types_added": standard_part_types,
        "category_attributes": len(category_attributes),
        "product_attributes": len(attributes),
        "references": len(references),
        "supplier_codes_kept_out_of_references": skipped_supplier_codes,
        "vehicles": len(vehicles),
        "fitments": len(fitments),
        "suppliers": len(suppliers_by_source_id),
        "offers": len(offers),
        "search_synonyms": len(synonyms),
        "supersessions": len(supersessions),
        "kits": len(kits),
        "media": len(media),
        "review_queue": len(review_items),
        "audits": len(audits),
        "raw_products": len(raw_products),
        "list_values": len(list_values),
        "field_docs": len(field_docs),
        "settings": len(settings),
    }


def import_workbook(session: Session, workbook_path: Path, replace: bool = False) -> dict[str, int]:
    if session.scalar(select(func.count()).select_from(Product)):
        if not replace:
            raise RuntimeError("Catalog is not empty; run with --replace for an atomic re-import")
        _clear_catalog(session)
        session.flush()

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "Products" in workbook.sheetnames:
        try:
            return _import_v2_workbook(session, workbook, workbook_path)
        finally:
            workbook.close()
    required = {
        "Master Products",
        "Reference Numbers",
        "Vehicles",
        "Fitments",
        "Supplier Offers",
        "Raw Mapping",
        "Cross Supplier Audit",
    }
    missing = required.difference(workbook.sheetnames)
    if missing:
        raise ValueError(f"Workbook is missing sheets: {', '.join(sorted(missing))}")

    master_rows = list(_rows(workbook["Master Products"]))
    reference_rows = list(_rows(workbook["Reference Numbers"]))
    vehicle_rows = list(_rows(workbook["Vehicles"]))
    fitment_rows = list(_rows(workbook["Fitments"]))
    offer_rows = list(_rows(workbook["Supplier Offers"]))
    raw_rows = list(_rows(workbook["Raw Mapping"]))
    audit_rows = list(_rows(workbook["Cross Supplier Audit"]))
    review_rows = (
        list(_rows(workbook["Review Queue"])) if "Review Queue" in workbook.sheetnames else []
    )
    workbook.close()

    supplier_names = {
        _text(row.get(key))
        for rows, key in (
            (reference_rows, "source_supplier"),
            (fitment_rows, "source_supplier"),
            (offer_rows, "supplier"),
            (raw_rows, "supplier"),
        )
        for row in rows
        if _text(row.get(key))
    }
    suppliers = {
        name: Supplier(
            name=name,
            normalized_name=normalize_text(name),
            created_at=datetime.now(UTC),
        )
        for name in sorted(supplier_names)
    }
    session.add_all(suppliers.values())

    brand_names = {_text(row.get("brand")) for row in master_rows if _text(row.get("brand"))}
    brands = {
        name: Brand(name=name, normalized_name=normalize_text(name)) for name in sorted(brand_names)
    }
    session.add_all(brands.values())
    session.flush()

    audit_by_product = {
        _text(row.get("product_id")): row for row in audit_rows if _text(row.get("product_id"))
    }
    products: dict[str, Product] = {}
    for row in master_rows:
        product_id = _text(row.get("product_id"))
        name = _text(row.get("product_name"))
        if not product_id or not name:
            raise ValueError(f"Master product is missing product_id or product_name: {row!r}")
        audit = audit_by_product.get(product_id, {})
        audit_status = (_text(audit.get("merge_status")) or "UNAUDITED").upper()
        product = Product(
            id=product_id,
            name=name,
            normalized_name=normalize_text(name),
            brand=brands.get(_text(row.get("brand"))),
            origin=_text(row.get("origin")),
            position=_text(row.get("position")),
            side=_text(row.get("side")),
            engine_hint=_text(row.get("engine_hint")),
            merge_basis=_text(row.get("merge_basis")),
            audit_status=audit_status,
            audit_reason=_text(audit.get("audit_reason")),
            is_searchable=audit_status != "CONFLICT",
        )
        products[product_id] = product
    session.add_all(products.values())

    vehicles: dict[str, Vehicle] = {}
    for row in vehicle_rows:
        vehicle_id = _text(row.get("vehicle_id"))
        model = _text(row.get("model"))
        if not vehicle_id or not model:
            continue
        make = _text(row.get("make"))
        generation = _text(row.get("generation"))
        vehicle = Vehicle(
            id=vehicle_id,
            make=make,
            model=model,
            generation=generation,
            normalized_text=normalize_text(" ".join(filter(None, [make, model, generation]))),
        )
        vehicles[vehicle_id] = vehicle
    session.add_all(vehicles.values())
    session.flush()

    searchable_fragments: dict[str, list[str]] = defaultdict(list)
    for product_id, product in products.items():
        searchable_fragments[product_id].extend(
            filter(
                None,
                [
                    product.name,
                    product.brand.name if product.brand else None,
                    product.position,
                    product.side,
                    product.engine_hint,
                ],
            )
        )

    reference_seen: set[tuple[str, str, str]] = set()
    references: list[ReferenceNumber] = []
    skipped_supplier_codes = 0
    for row in reference_rows:
        product_id = _text(row.get("product_id"))
        number = _text(row.get("number"))
        reference_type = (_text(row.get("type")) or "UNKNOWN").upper()
        if not product_id or product_id not in products or not number:
            continue
        if reference_type == "SUPPLIER_INTERNAL_CODE":
            skipped_supplier_codes += 1
            continue
        normalized_number = normalize_code(number)
        key = (product_id, normalized_number, reference_type)
        if not normalized_number or key in reference_seen:
            continue
        reference_seen.add(key)
        references.append(
            ReferenceNumber(
                product_id=product_id,
                number=number,
                normalized_number=normalized_number,
                type=reference_type,
                source_supplier=suppliers.get(_text(row.get("source_supplier"))),
            )
        )
        searchable_fragments[product_id].append(number)
    session.add_all(references)

    fitment_seen: set[tuple[str, str, str | None]] = set()
    fitments: list[Fitment] = []
    for row in fitment_rows:
        product_id = _text(row.get("product_id"))
        vehicle_id = _text(row.get("vehicle_id"))
        supplier_name = _text(row.get("source_supplier"))
        key = (product_id or "", vehicle_id or "", supplier_name)
        if (
            not product_id
            or product_id not in products
            or not vehicle_id
            or vehicle_id not in vehicles
            or key in fitment_seen
        ):
            continue
        fitment_seen.add(key)
        source_vehicle_text = _text(row.get("source_vehicle_text"))
        fitments.append(
            Fitment(
                product_id=product_id,
                vehicle_id=vehicle_id,
                source_vehicle_text=source_vehicle_text,
                source_supplier=suppliers.get(supplier_name),
            )
        )
        searchable_fragments[product_id].extend(
            filter(None, [vehicles[vehicle_id].normalized_text, source_vehicle_text])
        )
    session.add_all(fitments)

    offers: list[SupplierOffer] = []
    active_offer_keys: set[tuple[str, str, str]] = set()
    for row in offer_rows:
        product_id = _text(row.get("product_id"))
        supplier = suppliers.get(_text(row.get("supplier")))
        source_offer_id = _text(row.get("offer_id"))
        if not product_id or product_id not in products or not supplier or not source_offer_id:
            continue
        sku = _text(row.get("supplier_sku_or_ref"))
        normalized_sku = normalize_code(sku) or None
        offer_key = (supplier.name, product_id, normalized_sku or source_offer_id)
        is_active = offer_key not in active_offer_keys
        active_offer_keys.add(offer_key)
        offers.append(
            SupplierOffer(
                source_offer_id=source_offer_id,
                supplier=supplier,
                product_id=product_id,
                supplier_sku=sku,
                normalized_supplier_sku=normalized_sku,
                price=_decimal(row.get("demo_price_uzs")),
                currency="UZS",
                stock_qty=_integer(row.get("demo_stock")),
                updated_at=datetime.now(UTC),
                source_type="EXCEL",
                source_reference=f"{workbook_path.name}:{source_offer_id}",
                is_active=is_active,
                data_status="DEMO",
            )
        )
        if sku:
            searchable_fragments[product_id].append(sku)
    session.add_all(offers)

    review_by_raw = {
        _text(row.get("raw_id")): row for row in review_rows if _text(row.get("raw_id"))
    }
    raw_products: list[RawSupplierProduct] = []
    for row in raw_rows:
        raw_id = _text(row.get("raw_id"))
        product_id = _text(row.get("product_id"))
        supplier = suppliers.get(_text(row.get("supplier")))
        raw_name = _text(row.get("raw_name"))
        if not raw_id or not supplier or not raw_name:
            continue
        product_audit = products[product_id].audit_status if product_id in products else "UNAUDITED"
        review = review_by_raw.get(raw_id, {})
        raw_audit = (
            product_audit
            if product_audit == "CONFLICT"
            else (_text(review.get("status")) or product_audit)
        )
        audit_reason = (
            products[product_id].audit_reason
            if product_id in products and product_audit == "CONFLICT"
            else _text(review.get("review_reason"))
        )
        raw_products.append(
            RawSupplierProduct(
                source_raw_id=raw_id,
                supplier=supplier,
                source_row=_integer(row.get("source_row")),
                product_id=product_id if product_id in products else None,
                raw_name=raw_name,
                parsed_product_name=_text(row.get("parsed_product_name")),
                raw_brand=_text(row.get("brand")),
                supplier_code=_text(row.get("reference")),
                supplier_code_type=_text(row.get("reference_type")),
                origin=_text(row.get("origin")),
                unit=_text(row.get("unit")),
                vehicle_raw=_text(row.get("vehicle_raw")),
                position=_text(row.get("position")),
                side=_text(row.get("side")),
                engine_hint=_text(row.get("engine_hint")),
                audit_status=raw_audit,
                audit_reason=audit_reason,
                source_payload={key: value for key, value in row.items() if value is not None},
            )
        )
    session.add_all(raw_products)

    for product_id, product in products.items():
        product.search_text = normalize_text(" ".join(searchable_fragments[product_id]))

    session.flush()
    standard_part_types = ensure_standard_part_types(session)
    session.commit()
    return {
        "standard_part_types_added": standard_part_types,
        "products": len(products),
        "quarantined_conflicts": sum(not product.is_searchable for product in products.values()),
        "brands": len(brands),
        "references": len(references),
        "supplier_codes_kept_out_of_references": skipped_supplier_codes,
        "vehicles": len(vehicles),
        "fitments": len(fitments),
        "suppliers": len(suppliers),
        "offers": len(offers),
        "raw_products": len(raw_products),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Import the audited auto-parts XLSX catalog")
    parser.add_argument("workbook", type=Path)
    parser.add_argument(
        "--replace", action="store_true", help="Atomically replace existing catalog data"
    )
    args = parser.parse_args()
    if not args.workbook.is_file():
        parser.error(f"Workbook does not exist: {args.workbook}")

    with SessionLocal() as session:
        try:
            counts = import_workbook(session, args.workbook, replace=args.replace)
        except Exception:
            session.rollback()
            raise
    for name, count in counts.items():
        print(f"{name}: {count}")


if __name__ == "__main__":
    main()
