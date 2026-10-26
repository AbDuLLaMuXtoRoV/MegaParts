"""Read a seller's price list as it arrives: any .xls/.xlsx layout, headers or not.

Sellers send their own Excel exports. Column titles differ ("Наименование", "Модел",
"BREND", "YUAN", "NARXI") or are missing, headings between rows carry the warranty
("AMARTIZATOR GARANTIYA 10.000 KM 3 OY"), and the seller's phone numbers sit above the
table. The reader finds the columns by their titles and, where a column has none, by
what it holds, so a new seller's list needs no preparation.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from app.part_names import latin_upper

HEADER_ALIASES: dict[str, set[str]] = {
    "name": {"НАИМЕНОВАНИЕ", "НАИМЕНОВАНИЕ ТОВАРОВ", "НАИМЕНОВАНИЕ ТОВАРА", "НОМЕНКЛАТУРА",
             "ТОВАР", "NOMI", "NAME", "MAHSULOT", "TOVAR", "PRODUCT", "DESCRIPTION"},
    "model": {"МОДЕЛ", "МОДЕЛЬ", "MODEL", "MODEL AVTO", "AVTO", "MASHINA", "АВТО"},
    "brand": {"БРЕНД", "BREND", "BRAND", "МАРКА ДЕТАЛИ"},
    "unit": {"ЕД ИЗМ", "ЕДИЗМ", "ЕД.ИЗМ", "ЕД. ИЗМ", "UNIT", "O'LCHOV", "OLCHOV", "ЎЛЧОВ"},
    "country": {"ИШЛАБ ЧИКАРУВЧИ", "ИШЛАБ ЧИҚАРУВЧИ", "ISHLAB CHIQARUVCHI", "СТРАНА",
                "COUNTRY", "ПРОИЗВОДИТЕЛЬ", "DAVLAT"},
    "pack_qty": {"KAROBKA", "КОРОБКА", "KOROBKA", "УПАКОВКА", "QADOQ"},
    "price": {"NARXI", "NARX", "ЦЕНА", "PRICE", "НАРХ", "НАРХИ", "СУММА"},
    "price_cny": {"YUAN", "ЮАНЬ", "CNY"},
    "stock": {"QOLDIQ", "ОСТАТОК", "КОЛ-ВО", "КОЛИЧЕСТВО", "STOCK", "SONI", "КОЛ ВО"},
    "code": {"КОД", "KOD", "АРТИКУЛ", "ARTIKUL", "CODE", "SKU"},
    # A column the seller titles as the carmaker's number: always saved as an OEM code.
    "oem": {"OEM", "OE", "ОЕМ", "OEM KOD", "OEM KODI", "OEM NOMER", "OEM CODE", "OEM NUMBER",
            "OEM НОМЕР", "ОЕМ НОМЕР", "OEM КОД", "ОРИГИНАЛЬНЫЙ НОМЕР", "ОРИГИНАЛ НОМЕР",
            "КАТАЛОЖНЫЙ НОМЕР", "ZAVOD RAQAMI", "ORIGINAL NOMER"},
    "skip": {"№", "N", "NO", "ZAKAS", "ЗАКАЗ", "#"},
}
_HEADERS = {
    latin_upper(alias) or alias: column
    for column, aliases in HEADER_ALIASES.items()
    for alias in aliases
}
UNIT_WORDS = {"DONA", "KOMP", "SHT", "SH-T", "K-T", "KT", "KOMPLEKT", "PCS", "SET", "PAR", "JUFT"}
STATUS_WORDS = re.compile(r"YANGI KELDI|AKSIYA|NARXI? TUSHDI|YANGI NARX|SKIDKA|CHEGIRMA")
TOTAL_WORDS = re.compile(r"^(JAMI|ITOGO|VSEGO|TOTAL)\b")
_PHONE = re.compile(r"\+?998[\s-]?\(?\d{2}\)?[\s-]?\d{3}[\s-]?\d{2}[\s-]?\d{2}")
_DATE = re.compile(r"\b(\d{1,2})[.,/-](\d{1,2})[.,/-](20\d{2})\b")
# Car names that mark a "model" column; enough of them to tell it from part names.
_MODEL_HINT = re.compile(
    r"\b(NEXIA|NEKSIYA|DAMAS|COBALT|KOBALT|LACETTI|LASETTI|GENTRA|SPARK|MATIZ|TICO|TIKO|"
    r"CAPTIVA|MALIBU|EPICA|CIELO|ONIX|TRACKER|ORLANDO|CRUZE?|DOHC|SOHC|LABO|BEST|LANOS|"
    r"PORTER|BONGO|BYD|CHERY|ISUZU|EQUINOX|MONZA|AVEO)"
)


@dataclass
class PriceRow:
    row_number: int
    name: str
    code: str | None = None
    oem: str | None = None
    brand: str | None = None
    model: str | None = None
    unit: str | None = None
    country: str | None = None
    price: Decimal | None = None
    currency: str | None = None
    stock_qty: int | None = None
    pack_qty: int | None = None
    status: str | None = None
    section: str | None = None


@dataclass
class PriceList:
    rows: list[PriceRow] = field(default_factory=list)
    phones: list[str] = field(default_factory=list)
    telegram: str | None = None
    document_date: date | None = None
    columns: dict[str, int] = field(default_factory=dict)


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    return " ".join(str(value).split())


def read_sheet_rows(content: bytes, filename: str | None = None) -> list[list[str]]:
    """The rows of the workbook's fullest sheet as text; .xls needs the xlrd package."""
    sheets: list[list[list[str]]] = []
    if content[:4] == b"\xd0\xcf\x11\xe0" or (filename or "").lower().endswith(".xls"):
        import xlrd

        book = xlrd.open_workbook(file_contents=content)
        for sheet in book.sheets():
            sheets.append([[_cell(v) for v in sheet.row_values(i)] for i in range(sheet.nrows)])
    else:
        from openpyxl import load_workbook

        book = load_workbook(BytesIO(content), read_only=True, data_only=True)
        for sheet in book.worksheets:
            sheets.append([[_cell(v) for v in row] for row in sheet.iter_rows(values_only=True)])
        book.close()
    if not sheets:
        return []
    return max(sheets, key=lambda rows: sum(1 for row in rows if any(row)))


def _decimal(value: str) -> Decimal | None:
    cleaned = value.replace(" ", "").replace(",", ".")
    if not cleaned:
        return None
    try:
        number = Decimal(cleaned)
    except InvalidOperation:
        return None
    return number if number >= 0 else None


def _integer(value: str) -> int | None:
    number = _decimal(value)
    return int(number) if number is not None and number == int(number) else None


def _find_header(rows: list[list[str]]) -> tuple[int, dict[int, str]]:
    best = (-1, {})
    for index, row in enumerate(rows[:20]):
        found = {}
        for column, value in enumerate(row):
            if not value:
                continue
            key = _HEADERS.get(latin_upper(value).rstrip(":.") or value)
            if key and key not in found.values():
                found[column] = key
        if len(found) >= 2 and len(found) > len(best[1]):
            best = (index, found)
    return best


def _share(values: list[str], test) -> float:
    return sum(1 for value in values if test(value)) / len(values)


def _field(row: list[str], columns: dict[str, int], name: str) -> str:
    column = columns.get(name)
    return row[column] if column is not None and column < len(row) else ""


def _infer_columns(rows: list[list[str]], known: dict[int, str]) -> dict[int, str]:
    """Name the untitled columns by what most of their cells look like."""
    columns = dict(known)
    width = max((len(row) for row in rows), default=0)
    stats: dict[int, dict[str, float]] = {}
    for column in range(width):
        if column in columns:
            continue
        values = [row[column] for row in rows if column < len(row) and row[column]]
        # Columns filled in under a tenth of the rows are notes, not data.
        if not values or len(values) < len(rows) * 0.1:
            continue
        upper = [latin_upper(value) for value in values]
        stats[column] = {
            "unit": _share(upper, lambda v: v.replace(".", "") in UNIT_WORDS),
            "model": _share(upper, lambda v: bool(_MODEL_HINT.search(v))),
            "status": _share(upper, lambda v: bool(STATUS_WORDS.search(v))),
            "number": _share(upper, lambda v: _decimal(v) is not None),
            "code": _share(upper, lambda v: sum(c.isdigit() for c in v) >= 4 and len(v.split()) <= 3),
            "length": sum(len(v) for v in values) / len(values),
            "filled": len(values) / len(rows),
        }
    taken = set(columns.values())

    def claim(field_name: str, test) -> None:
        if field_name in taken:
            return
        candidates = [column for column, stat in stats.items() if column not in columns and test(stat)]
        if candidates:
            column = max(candidates, key=lambda c: stats[c]["filled"])
            columns[column] = field_name
            taken.add(field_name)

    claim("unit", lambda s: s["unit"] > 0.6)
    claim("status", lambda s: s["status"] > 0.6)
    claim("model", lambda s: s["model"] > 0.4 and s["length"] < 30)
    # Names are the longest text; a code column ("96553241-N/WSM") is shorter.
    if "name" not in taken:
        text_columns = [
            column for column, stat in stats.items()
            if column not in columns and stat["number"] < 0.3
        ]
        if text_columns:
            column = max(text_columns, key=lambda c: stats[c]["length"] * stats[c]["filled"])
            columns[column] = "name"
            taken.add("name")
    claim("code", lambda s: s["code"] > 0.5 and s["number"] < 0.5 and s["length"] < 25)
    return columns


def _status(row: list[str], columns: dict[int, str]) -> str | None:
    for column, value in enumerate(row):
        if columns.get(column) in (None, "status") and STATUS_WORDS.search(latin_upper(value)):
            return value
    return None


def parse_price_list(content: bytes, filename: str | None = None) -> PriceList:
    rows = read_sheet_rows(content, filename)
    result = PriceList()
    header_index, known = _find_header(rows)
    body = rows[header_index + 1:] if header_index >= 0 else rows
    columns = _infer_columns(body, known)
    result.columns = {name: column for column, name in columns.items()}
    if "name" not in result.columns:
        raise ValueError("Could not find the column with part names")

    # The seller's phones and the list's date usually sit above the table.
    for row in rows[: max(header_index, 0) + 1] + rows[-3:]:
        text = " ".join(row)
        for match in _PHONE.finditer(text):
            phone = re.sub(r"[^\d+]", "", match.group())
            phone = phone if phone.startswith("+") else f"+{phone}"
            if phone not in result.phones:
                result.phones.append(phone)
            after = text[match.end(): match.end() + 15].upper()
            if "TELEGRAM" in after:
                result.telegram = phone
    for text in [filename or "", *(" ".join(row) for row in rows[:6])]:
        if match := _DATE.search(text):
            day, month, year = map(int, match.groups())
            try:
                result.document_date = date(year, month, day)
            except ValueError:
                continue
            break

    price_values: list[Decimal] = []
    section = None
    name_column = result.columns["name"]
    for offset, row in enumerate(body):
        row_number = header_index + offset + 2
        filled = [value for value in row if value]
        if not filled:
            continue
        name = row[name_column] if name_column < len(row) else ""
        if len(filled) == 1 and not _decimal(filled[0] or ""):
            # A heading between rows ("PORSHIN GARANTIYA 30.000 KM 3 OY") applies below it.
            if not TOTAL_WORDS.search(latin_upper(filled[0])):
                section = filled[0]
            continue
        if not name or TOTAL_WORDS.search(latin_upper(name)) or _decimal(name) is not None:
            continue

        def value(field_name: str, row: list[str] = row) -> str:
            return _field(row, result.columns, field_name)

        price = _decimal(value("price"))
        if price is not None:
            price_values.append(price)
        result.rows.append(
            PriceRow(
                row_number=row_number,
                name=name,
                code=value("code") or None,
                oem=value("oem") or None,
                brand=value("brand") or None,
                model=value("model") or None,
                unit=value("unit") or None,
                country=value("country") or None,
                price=price,
                stock_qty=_integer(value("stock")),
                pack_qty=_integer(value("pack_qty")),
                status=_status(row, columns),
                section=section,
            )
        )
    # Prices in so'm run to thousands; a list whose prices stay small is in dollars.
    currency = "USD" if price_values and max(price_values) < 10_000 else "UZS"
    for row in result.rows:
        if row.price is not None:
            row.price = row.price.quantize(Decimal("0.01"))
            row.currency = currency
    return result


_WARRANTY_KM = re.compile(r"(\d{1,3}(?:[.,\s]\d{3})+|\d+)\s*KM")
_WARRANTY_MONTHS = re.compile(r"(\d+)\s*OY\b")


def warranty(section: str | None) -> tuple[int | None, int | None]:
    """(months, km) from a heading such as "GARANTIYA 10.000 KM 3 OY"."""
    text = latin_upper(section or "")
    if "GARANT" not in text and "KAFOLAT" not in text:
        return None, None
    km = _WARRANTY_KM.search(text)
    months = _WARRANTY_MONTHS.search(text)
    return (
        int(months.group(1)) if months else None,
        int(re.sub(r"\D", "", km.group(1))) if km else None,
    )


def promo(status: str | None) -> tuple[str | None, int | None]:
    """(label, discount percent) from "YANGI KELDI", "-20% AKSIYA", "NARXI TUSHDI"."""
    text = latin_upper(status or "")
    if not text:
        return None, None
    if match := re.search(r"(\d{1,2})\s*%", status or ""):
        return "DISCOUNT", int(match.group(1))
    if "AKSIYA" in text or "SKIDKA" in text or "CHEGIRMA" in text:
        return "DISCOUNT", None
    if "TUSHDI" in text:
        return "PRICE_DROP", None
    if "YANGI NARX" in text:
        return "NEW_PRICE", None
    if "YANGI KELDI" in text:
        return "NEW_ARRIVAL", None
    return None, None
