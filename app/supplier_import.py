from __future__ import annotations

import re
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (
    Brand,
    OfferImportReview,
    Product,
    ReferenceNumber,
    Supplier,
    SupplierImport,
    SupplierOffer,
    SupplierProductMapping,
)
from app.normalization import normalize_code, normalize_text

HEADER_ALIASES = {
    "supplier_sku": {"supplier sku", "sku", "supplier code", "article", "code"},
    "reference": {"reference", "reference number", "oem", "oem number"},
    "brand": {"brand", "manufacturer"},
    "product_name": {"product", "product name", "name", "description"},
    "price": {"price", "narx", "price uzs"},
    "currency": {"currency", "valyuta"},
    "stock_qty": {"stock", "stock qty", "quantity", "qty", "qoldiq"},
    "warehouse": {"warehouse", "ombor"},
    "updated_at": {"updated at", "updated", "date", "timestamp"},
    "part_origin": {"origin", "part origin", "turi", "тип", "происхождение"},
    "condition": {"condition", "holati", "состояние"},
    "warranty_months": {"warranty", "warranty months", "kafolat", "гарантия"},
    "return_days": {"returns", "return days", "qaytarish", "возврат"},
    "lead_time_days": {
        "lead time",
        "lead days",
        "delivery days",
        "muddat",
        "yetkazish muddati",
        "срок",
        "срок поставки",
    },
}

# Offer labels as suppliers write them; anything else is left blank rather than rejected.
ORIGIN_VALUES = {
    "GENUINE": {"genuine", "original", "oem", "gm", "original gm", "asl", "оригинал"},
    "OE": {"oe", "oe brand", "oes"},
    "AFTERMARKET": {"aftermarket", "analog", "analogue", "replacement", "аналог"},
    "BUDGET": {"budget", "economy", "arzon", "эконом"},
}
CONDITION_VALUES = {
    "NEW": {"new", "yangi", "новый", "новая", "новое"},
    "USED": {"used", "ishlatilgan", "бу", "б у"},
    "REMANUFACTURED": {"reman", "remanufactured", "rebuilt", "восстановленный"},
}
COUNT_FIELDS = ("warranty_months", "return_days", "lead_time_days")


def _header(value: Any) -> str:
    return normalize_text(str(value or "").replace("_", " ").replace("-", " "))


def _field_map(headers: list[Any]) -> dict[int, str]:
    result: dict[int, str] = {}
    for index, raw in enumerate(headers):
        normalized = _header(raw)
        for field, aliases in HEADER_ALIASES.items():
            if normalized in aliases:
                result[index] = field
                break
    return result


def _text(value: Any) -> str | None:
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def _decimal(value: Any) -> Decimal | None:
    if value in (None, ""):
        return None
    cleaned = str(value).strip().replace(" ", "").replace(",", ".")
    try:
        return Decimal(cleaned)
    except InvalidOperation as exc:
        raise ValueError(f"Invalid price: {value}") from exc


def _integer(value: Any) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(float(str(value).strip().replace(",", ".")))
    except ValueError as exc:
        raise ValueError(f"Invalid stock quantity: {value}") from exc


def _choice(value: Any, choices: dict[str, set[str]]) -> str | None:
    words = normalize_text(value)
    return next((code for code, spellings in choices.items() if words in spellings), None)


def _count(value: Any) -> int | None:
    """Days or months written as "12", "12 oy" or "14 дней"."""
    match = re.search(r"\d+", str(value or ""))
    return int(match.group()) if match else None


def _timestamp(value: Any) -> datetime:
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=UTC)
    if value not in (None, ""):
        try:
            parsed = datetime.fromisoformat(str(value).strip())
            return parsed if parsed.tzinfo else parsed.replace(tzinfo=UTC)
        except ValueError:
            pass
    return datetime.now(UTC)


def _json_payload(row: dict[str, Any]) -> dict[str, Any]:
    return {
        key: value.isoformat() if isinstance(value, datetime) else str(value)
        for key, value in row.items()
        if value not in (None, "")
    }


def _match_product(
    session: Session, supplier_id: int, row: dict[str, Any]
) -> tuple[Product | None, str | None]:
    sku = _text(row.get("supplier_sku"))
    normalized_sku = normalize_code(sku)
    if normalized_sku:
        mapping = session.scalar(
            select(SupplierProductMapping).where(
                SupplierProductMapping.supplier_id == supplier_id,
                SupplierProductMapping.normalized_supplier_sku == normalized_sku,
                SupplierProductMapping.is_confirmed.is_(True),
                SupplierProductMapping.mapping_type.in_(("SUPPLIER_SKU", "EXACT")),
            )
        )
        if mapping:
            return session.get(Product, mapping.product_id), "SUPPLIER_SKU_MAPPING"

        existing = session.scalar(
            select(SupplierOffer).where(
                SupplierOffer.supplier_id == supplier_id,
                SupplierOffer.normalized_supplier_sku == normalized_sku,
                SupplierOffer.is_active.is_(True),
            )
        )
        if existing:
            return session.get(Product, existing.product_id), "EXISTING_OFFER"

    brand_name = normalize_text(_text(row.get("brand")))
    reference = normalize_code(_text(row.get("reference")))
    if brand_name and reference:
        products = (
            session.scalars(
                select(Product)
                .join(Brand, Brand.id == Product.brand_id)
                # Explicit: a bare join would follow the number's own brand, not the product.
                .join(ReferenceNumber, ReferenceNumber.product_id == Product.id)
                .where(
                    Brand.normalized_name == brand_name,
                    ReferenceNumber.normalized_number == reference,
                    Product.is_searchable.is_(True),
                )
            )
            .unique()
            .all()
        )
        if len(products) == 1:
            return products[0], "BRAND_REFERENCE"
    if normalized_sku:
        mapping = session.scalar(
            select(SupplierProductMapping)
            .where(
                SupplierProductMapping.supplier_id == supplier_id,
                SupplierProductMapping.normalized_supplier_sku == normalized_sku,
                SupplierProductMapping.is_confirmed.is_(True),
                SupplierProductMapping.mapping_type.in_(("CONFIRMED", "MANUAL")),
            )
            .order_by(
                SupplierProductMapping.mapping_type == "MANUAL",
                SupplierProductMapping.id,
            )
        )
        if mapping:
            return session.get(Product, mapping.product_id), mapping.mapping_type
    return None, None


def import_supplier_workbook(
    session: Session,
    supplier: Supplier,
    content: bytes,
    filename: str | None = None,
) -> SupplierImport:
    batch = SupplierImport(
        supplier_id=supplier.id,
        source_filename=filename,
        source_type="EXCEL",
        status="PROCESSING",
    )
    session.add(batch)
    session.flush()

    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    try:
        headers = list(next(rows))
    except StopIteration as exc:
        workbook.close()
        batch.status = "FAILED"
        batch.errors = 1
        raise ValueError("The workbook is empty") from exc
    fields = _field_map(headers)
    if "supplier_sku" not in fields.values():
        workbook.close()
        batch.status = "FAILED"
        batch.errors = 1
        raise ValueError("A supplier SKU/code column is required")

    for row_number, values in enumerate(rows, start=2):
        row = {field: values[index] for index, field in fields.items()}
        if not any(value not in (None, "") for value in row.values()):
            continue
        batch.rows_received += 1
        sku = _text(row.get("supplier_sku"))
        normalized_sku = normalize_code(sku)
        try:
            if not sku or not normalized_sku:
                raise ValueError("Supplier SKU is required")
            price = _decimal(row.get("price"))
            stock_qty = _integer(row.get("stock_qty"))
            product, match_type = _match_product(session, supplier.id, row)
            if product is None:
                batch.unmatched_rows += 1
                session.add(
                    OfferImportReview(
                        import_id=batch.id,
                        supplier_id=supplier.id,
                        supplier_sku=sku,
                        raw_payload={"row": row_number, **_json_payload(row)},
                    )
                )
                continue

            batch.rows_matched += 1
            offer = session.scalar(
                select(SupplierOffer).where(
                    SupplierOffer.supplier_id == supplier.id,
                    SupplierOffer.product_id == product.id,
                    SupplierOffer.normalized_supplier_sku == normalized_sku,
                    SupplierOffer.is_active.is_(True),
                )
            )
            if offer is None:
                offer = SupplierOffer(
                    source_offer_id=uuid4().hex,
                    supplier_id=supplier.id,
                    product_id=product.id,
                    supplier_sku=sku,
                    normalized_supplier_sku=normalized_sku,
                )
                session.add(offer)
                batch.new_offers += 1
            else:
                batch.updated_offers += 1
            offer.price = price
            offer.currency = (_text(row.get("currency")) or "UZS").upper()[:3]
            offer.stock_qty = stock_qty
            offer.warehouse = _text(row.get("warehouse"))
            offer.updated_at = _timestamp(row.get("updated_at"))
            offer.source_type = "EXCEL"
            offer.source_reference = f"import:{batch.id}:row:{row_number}:{match_type}"
            offer.data_status = "LIVE"
            offer.is_active = True
            # Optional columns only overwrite what the file actually carries.
            present = set(fields.values())
            if "part_origin" in present:
                offer.part_origin = _choice(row.get("part_origin"), ORIGIN_VALUES)
            if "condition" in present:
                offer.condition = _choice(row.get("condition"), CONDITION_VALUES)
            for field in COUNT_FIELDS:
                if field in present:
                    setattr(offer, field, _count(row.get(field)))
        except (ValueError, TypeError) as exc:
            batch.errors += 1
            session.add(
                OfferImportReview(
                    import_id=batch.id,
                    supplier_id=supplier.id,
                    supplier_sku=sku,
                    raw_payload={"row": row_number, **_json_payload(row)},
                    status="ERROR",
                    error=str(exc),
                )
            )

    workbook.close()
    batch.status = "COMPLETED_WITH_ERRORS" if batch.errors else "COMPLETED"
    session.flush()
    return batch
